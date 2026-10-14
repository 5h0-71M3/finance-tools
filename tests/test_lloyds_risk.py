from datetime import datetime

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from finance_tools.lloyds_risk.build import CANADA, UK, build_region, read_frame, write_deliverable
from finance_tools.lloyds_risk.server import find_old_files

UK_UMR = "Lloyd's UMR Code"

_LOOKER_SPELLING = {  # how the Canada look names the deliverable's fields
    "Coalition Policy Number": "Quote Coalition Policy Number", "Effective Date": "Quote Effective Date Date",
    "Total Bound Premium": "Bound Premium Incl Endorsements", "Lloyds Umr Code": "Lloyd's UMR Code",
}
_UK_AS_CANADA = {"Lloyd's UMR Code": "Lloyds Umr Code", "Lloyd's Risk Share": "Lloyds Risk Share",
                 "Company Full Address": "Quote Company Address Complete", "Lifecycle State": "Quote Status"}


def _new(spec, rows: list[tuple[str, str]], headers: list[str] | None = None) -> pd.DataFrame:
    """a Looker download with (policy, looker UMR) rows; every other cell is blank"""
    headers = headers or spec.output_headers[int(spec.duplicate_key):]
    df = pd.DataFrame(np.nan, index=range(len(rows)), columns=headers, dtype=object)
    umr_col = next(c for c in headers if "umr" in c.lower())
    df[headers[0]] = [p for p, _ in rows]
    df[umr_col] = [u for _, u in rows]
    df[headers[1]] = datetime(2026, 9, 30)
    return df


def _old(rows: list[tuple[str, str]]) -> pd.DataFrame:
    return pd.DataFrame({"Coalition Policy Number": [p for p, _ in rows], "Lloyds Umr Code": [u for _, u in rows]})


def test_umr_priority_is_old_24_then_old_25_then_looker():
    new = _new(UK, [("P1", UK.umr_25), ("P2", UK.umr_24), ("P3", UK.umr_24), ("P4", UK.umr_24)])
    old_24 = _old([("P1", UK.umr_24)])
    old_25 = _old([("P1", UK.umr_25), ("P2", UK.umr_25)])

    r = build_region(new, old_24, old_25, UK)

    assert list(r.file_24["Coalition Policy Number"]) == ["P1", "P3", "P4"]
    assert list(r.file_25["Coalition Policy Number"]) == ["P2"]
    assert set(r.file_24[UK_UMR]) == {UK.umr_24} and set(r.file_25[UK_UMR]) == {UK.umr_25}
    assert (r.from_old_24, r.from_old_25, r.from_looker) == (1, 1, 2)
    assert r.umr_changed_by_carry_forward == 2


def test_blanks_become_zero_and_order_is_kept():
    new = _new(UK, [("P2", UK.umr_24), ("P1", UK.umr_24)])
    r = build_region(new, _old([]), _old([]), UK)

    assert list(r.file_24["Coalition Policy Number"]) == ["P2", "P1"]
    assert (r.file_24["Pollution Sublimit"] == 0).all()
    assert r.file_24["Effective Date"].iloc[0] == pd.Timestamp(2026, 9, 30)


def test_unexpected_umr_goes_to_25_file_and_is_reported():
    new = _new(UK, [("P1", UK.umr_25), ("P2", UK.umr_25)])
    old_25 = _old([("P1", "B174010161SS24S")])

    r = build_region(new, _old([]), old_25, UK)

    assert r.file_24.empty
    assert list(r.file_25[UK_UMR]) == ["B174010161SS24S", UK.umr_25]
    assert r.unexpected_umrs == {"B174010161SS24S": 1}


def test_canada_renames_headers_and_duplicates_policy_number():
    looker_headers = [_LOOKER_SPELLING.get(h, h) for h in CANADA.field_headers]
    new = _new(CANADA, [("C1", CANADA.umr_25)], headers=looker_headers)
    old_24 = pd.DataFrame({"Coalition Policy Number": ["C1"], "Coalition Policy Number.1": ["C1"],
                           "Lloyds Umr Code": [CANADA.umr_24]})

    r = build_region(new, old_24, _old([]), CANADA)

    assert list(r.file_24.columns) == CANADA.output_headers
    assert list(r.file_24.iloc[0, :2]) == ["C1", "C1"]
    assert r.file_24["Lloyds Umr Code"].iloc[0] == CANADA.umr_24


def test_reordered_looker_columns_are_matched_by_name():
    # Sept 2026: the UK look came back with Canada-style names and State moved to the end
    reordered = [h for h in UK.field_headers if h != "Quote Company State"] + ["Quote Company State"]
    new = _new(UK, [("P1", UK.umr_24)], headers=[_UK_AS_CANADA.get(h, h) for h in reordered])
    new["Quote Company State"] = "ENG"
    new["Lloyds Risk Share"] = 0.2

    r = build_region(new, _old([]), _old([]), UK)

    assert list(r.file_24.columns) == UK.output_headers
    assert r.file_24["Quote Company State"].iloc[0] == "ENG"
    assert r.file_24["Lloyd's Risk Share"].iloc[0] == 0.2
    assert r.file_24[UK_UMR].iloc[0] == UK.umr_24


def test_wrong_file_is_rejected():
    with pytest.raises(ValueError, match="Unrecognised: \\['a', 'b', 'c'\\]"):
        build_region(pd.DataFrame(columns=["a", "b", "c"]), _old([]), _old([]), UK)


def test_swapped_region_is_rejected():
    canada_download = _new(UK, [("C1", CANADA.umr_24), ("C2", CANADA.umr_25)])  # same shape, Canada UMRs
    with pytest.raises(ValueError, match="doesn't look like a UK file"):
        build_region(canada_download, _old([]), _old([]), UK)

    uk_download = _new(UK, [("P1", UK.umr_24)])
    with pytest.raises(ValueError, match=f"Last month's {UK.umr_25} doesn't look like a UK file"):
        build_region(uk_download, _old([]), _old([("C1", CANADA.umr_25)]), UK)


def test_written_file_has_sheet_and_date_format(tmp_path):
    r = build_region(_new(UK, [("P1", UK.umr_24)]), _old([]), _old([]), UK)
    path = write_deliverable(r.file_24, tmp_path / "out.xlsx", UK.sheet_24)

    ws = load_workbook(path)[UK.sheet_24]
    assert ws["A2"].value == "P1"
    assert ws["B2"].number_format == "mm-dd-yy"


def test_find_old_files_does_not_confuse_ss24_with_ss24s(tmp_path):
    (tmp_path / "B174010161SS24S Risk BDX 2026-09-04T1148.xlsx").touch()
    (tmp_path / "B174010161SS24 Risk BDX 2026-09-04T1148.xlsx").touch()

    found = find_old_files(tmp_path)

    assert found["B174010161SS24"].name.startswith("B174010161SS24 ")
    assert found["B174010160SS25"] is None


def test_read_frame_drops_looker_row_numbers(tmp_path):
    path = tmp_path / "looker.csv"
    pd.DataFrame({"Unnamed: 0": [1, 2], "Coalition Policy Number": ["P1", "P2"]}).to_csv(path, index=False)

    assert list(read_frame(path).columns) == ["Coalition Policy Number"]
